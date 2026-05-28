#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Генерация инструкции по деплою в формате Word (.docx)."""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BLUE = RGBColor(0x1B, 0x3A, 0x8B)
GOLD = RGBColor(0xA8, 0x82, 0x3C)
DARK = RGBColor(0x1A, 0x1A, 0x2E)
GRAY = RGBColor(0x55, 0x55, 0x60)
CODE_BG = "F2F1ED"
HEAD_BG = "1B3A8B"
ZEBRA_BG = "F3ECE0"

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "ERA-Coffee-Деплой-инструкция.docx")

doc = Document()

# базовый шрифт
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.font.color.rgb = DARK
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15


def shade(paragraph, fill):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def set_cell_bg(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def h1(text):
    p = doc.add_paragraph()
    p.space_before = Pt(18)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = BLUE
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(10)


def h2(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = BLUE
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    # тонкая линия под заголовком
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "C9A876")
    pbdr.append(bottom)
    pPr.append(pbdr)


def para(text, runs=None):
    p = doc.add_paragraph()
    if runs:
        for t, bold in runs:
            r = p.add_run(t)
            r.bold = bold
    else:
        p.add_run(text)
    return p


def bullet(text, checklist=False):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(("☐  " if checklist else "") + text)
    return p


def numbered(text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p


def code(lines):
    p = doc.add_paragraph()
    shade(p, CODE_BG)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Pt(8)
    p.paragraph_format.right_indent = Pt(8)
    for i, line in enumerate(lines):
        r = p.add_run(line)
        r.font.name = "Consolas"
        rpr = r._element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        rfonts.set(qn("w:ascii"), "Consolas")
        rfonts.set(qn("w:hAnsi"), "Consolas")
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0x22, 0x22, 0x2A)
        if i < len(lines) - 1:
            r.add_break()


def note(text):
    p = doc.add_paragraph()
    shade(p, "FBF3E2")
    p.paragraph_format.left_indent = Pt(8)
    p.paragraph_format.right_indent = Pt(8)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("⚠  " + text)
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0x6B, 0x4E, 0x1E)


def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        set_cell_bg(hdr[i], HEAD_BG)
        para_c = hdr[i].paragraphs[0]
        run = para_c.add_run(htext)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(9.5)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci, val in enumerate(row):
            if ri % 2 == 1:
                set_cell_bg(cells[ci], ZEBRA_BG)
            rp = cells[ci].paragraphs[0]
            run = rp.add_run(val)
            run.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


# ─────────────────────────── ТИТУЛ ───────────────────────────
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = title.add_run("ERA Coffee")
r.bold = True
r.font.size = Pt(30)
r.font.color.rgb = BLUE
sub = doc.add_paragraph()
r = sub.add_run("Инструкция по размещению сайта на Reg.ru (домен + VPS)")
r.font.size = Pt(13)
r.font.color.rgb = GRAY
meta = doc.add_paragraph()
r = meta.add_run("Домен eracoffee.ru · Node.js + nginx + PM2 + бесплатный SSL · почта для заявок")
r.font.size = Pt(10)
r.italic = True
r.font.color.rgb = GRAY

para("")
para(
    "Полная пошаговая инструкция. Все команды можно копировать как есть. "
    "Места, где нужно подставить свои данные, помечены ЗАГЛАВНЫМИ_БУКВАМИ "
    "(например, ВАШ_IP_АДРЕС)."
)

note(
    "Про формы: доставка заявок на почту уже реализована в коде. Чтобы письма "
    "приходили — задайте SMTP-доступ вашего ящика в .env.local (Шаг 11). Пока SMTP "
    "не настроен, заявки дублируются в лог сервера (pm2 logs) и не теряются."
)

h2("Что в итоге получится")
code([
    "Посетитель → eracoffee.ru (HTTPS)",
    "           → nginx (порты 80/443, SSL Let's Encrypt)",
    "           → Next.js (next start, порт 3000, под управлением PM2)",
])
para("Ориентировочная стоимость: домен .ru ~200–300 ₽/год + VPS ~400–700 ₽/мес.")

# ─────────────────────────── ШАГИ ───────────────────────────
h1("Шаг 1. Регистрация домена eracoffee.ru")
numbered("Зайдите на reg.ru, создайте аккаунт или войдите.")
numbered("В поиске доменов введите eracoffee.ru, добавьте в корзину, оплатите.")
numbered("После оплаты домен появится в разделе «Мои домены».")
numbered("Пока ничего в DNS не меняйте — вернёмся к этому в Шаге 8, когда будет готов VPS и известен его IP.")

h1("Шаг 2. Заказ VPS")
numbered("На Reg.ru: Услуги → VPS/VDS → Заказать.")
para("Рекомендуемая конфигурация (минимум для сборки Next.js):")
bullet("CPU: 2 ядра")
bullet("RAM: 2 ГБ (лучше 4 ГБ — сборка Next.js прожорлива по памяти)")
bullet("Диск: 20–30 ГБ SSD")
bullet("ОС: Ubuntu 22.04 LTS (инструкция написана под неё)")
numbered("После создания придут: IP-адрес сервера, логин root и пароль (или вы зададите SSH-ключ).")
numbered("Запишите IP — он понадобится в Шаге 8.")

h1("Шаг 3. Первое подключение по SSH")
para("На Mac откройте Терминал (на Windows — PowerShell):")
code(["ssh root@ВАШ_IP_АДРЕС"])
para("Введите пароль (символы не отображаются — это нормально). При первом входе подтвердите yes.")
para("Обновите систему и поставьте базовые утилиты:")
code(["apt update && apt upgrade -y", "apt install -y git curl ufw"])
para("Настройте файрвол (откроем SSH и web-порты):")
code(["ufw allow OpenSSH", "ufw allow 'Nginx Full'", "ufw --force enable"])
note("'Nginx Full' станет доступен после установки nginx (Шаг 5). Если ufw ругается — выполните строку ещё раз после Шага 5.")

h1("Шаг 4. Установка Node.js 20 LTS")
code([
    "curl -fsSL https://deb.nodesource.com/setup_20.x | bash -",
    "apt install -y nodejs",
    "node -v   # должно показать v20.x",
    "npm -v",
])
para("Установите PM2 (держит сайт запущенным и перезапускает после ребута):")
code(["npm install -g pm2"])

h1("Шаг 5. Установка nginx")
code(["apt install -y nginx", "systemctl enable nginx", "systemctl start nginx"])
para("Проверьте: откройте в браузере http://ВАШ_IP — должна быть страница-заглушка nginx.")
para("Если на Шаге 3 ufw не дал добавить 'Nginx Full' — сделайте сейчас:")
code(["ufw allow 'Nginx Full'"])

h1("Шаг 6. Клонирование проекта и сборка")
para("Репозиторий публичный, логин не нужен.")
code([
    "cd /var/www",
    "git clone https://github.com/1314johnsmith1314-cloud/era-coffee.git",
    "cd era-coffee",
    "npm ci          # установка зависимостей по package-lock",
    "npm run build   # production-сборка",
])
note("Если сборка падает с ошибкой памяти (Killed / JS heap out of memory) на VPS с 2 ГБ RAM — добавьте swap и повторите npm run build (команды ниже).")
code([
    "fallocate -l 2G /swapfile && chmod 600 /swapfile",
    "mkswap /swapfile && swapon /swapfile",
    "echo '/swapfile none swap sw 0 0' >> /etc/fstab",
])

h1("Шаг 7. Запуск через PM2")
code([
    "cd /var/www/era-coffee",
    "pm2 start npm --name era-coffee -- start",
    "pm2 save",
    "pm2 startup systemd      # выполните команду, которую PM2 распечатает",
])
para("Проверьте, что сайт отвечает локально на сервере:")
code(["curl -I http://localhost:3000"])
para("Должно прийти HTTP/1.1 200 OK.")
para("Полезные команды PM2 на будущее:")
bullet("pm2 logs era-coffee — логи (в т.ч. заявки с форм)")
bullet("pm2 restart era-coffee — перезапуск")
bullet("pm2 status — статус")

h1("Шаг 8. Привязка домена (DNS в Reg.ru)")
numbered("Reg.ru → Мои домены → eracoffee.ru → Управление зоной DNS.")
numbered("Убедитесь, что используются DNS-серверы Reg.ru (ns1.reg.ru, ns2.reg.ru). Если стоят чужие — переключите на Reg.ru.")
numbered("Удалите дефолтные A/CNAME-записи на парковку и добавьте:")
table(["Тип", "Имя (host)", "Значение"], [
    ["A", "@", "ВАШ_IP_АДРЕС"],
    ["A", "www", "ВАШ_IP_АДРЕС"],
])
numbered("Сохраните. Распространение DNS — от 15 минут до нескольких часов.")
para("Проверить, что домен указывает на ваш сервер:")
code(["ping eracoffee.ru        # должен отвечать ваш IP"])

h1("Шаг 9. Настройка nginx (reverse proxy)")
para("Создайте конфиг сайта:")
code(["nano /etc/nginx/sites-available/eracoffee.ru"])
para("Вставьте:")
code([
    "server {",
    "    listen 80;",
    "    server_name eracoffee.ru www.eracoffee.ru;",
    "",
    "    location /_next/static/ {",
    "        alias /var/www/era-coffee/.next/static/;",
    "        expires 1y;",
    '        add_header Cache-Control "public, immutable";',
    "    }",
    "",
    "    location / {",
    "        proxy_pass http://localhost:3000;",
    "        proxy_http_version 1.1;",
    "        proxy_set_header Upgrade $http_upgrade;",
    "        proxy_set_header Connection 'upgrade';",
    "        proxy_set_header Host $host;",
    "        proxy_set_header X-Real-IP $remote_addr;",
    "        proxy_set_header X-Forwarded-For $proxy_add_x_forwarded_for;",
    "        proxy_set_header X-Forwarded-Proto $scheme;",
    "        proxy_cache_bypass $http_upgrade;",
    "    }",
    "",
    "    client_max_body_size 25M;",
    "}",
])
para("Сохраните (Ctrl+O, Enter, Ctrl+X). Активируйте конфиг:")
code([
    "ln -s /etc/nginx/sites-available/eracoffee.ru /etc/nginx/sites-enabled/",
    "rm -f /etc/nginx/sites-enabled/default",
    "nginx -t        # проверка синтаксиса — должно быть ok и successful",
    "systemctl reload nginx",
])
para("Теперь http://eracoffee.ru должен открывать сайт (пока без HTTPS).")

h1("Шаг 10. Бесплатный SSL (HTTPS) через Let's Encrypt")
code([
    "apt install -y certbot python3-certbot-nginx",
    "certbot --nginx -d eracoffee.ru -d www.eracoffee.ru",
])
para("Certbot спросит email, согласие с условиями и предложит автоматически перенаправлять HTTP → HTTPS — выберите да (2).")
para("Сертификат продлевается автоматически. Проверить автопродление:")
code(["certbot renew --dry-run"])
para("Готово — сайт доступен по https://eracoffee.ru.")

h1("Шаг 11. Доставка заявок с форм на почту (ВАЖНО)")
para(
    "Код доставки уже написан (app/api/contact/route.ts) — нужно только задать "
    "переменные окружения. Отправка идёт через ваш собственный почтовый ящик по SMTP, "
    "без сторонних сервисов. Если переменные не заданы — заявка всё равно пишется в "
    "pm2 logs как резервная копия."
)
para("Заявка приходит письмом с темой «ERA Coffee — <тип формы>» и списком полей (имя, телефон, тип бизнеса, выбранные сорта, объём, цена и т.д.).")
h2("Настройка SMTP (на примере Яндекс 360)")
para("Сначала заведите ящик info@eracoffee.ru (см. Шаг 13), затем:")
numbered("В Яндекс ID → Безопасность → Пароли приложений создайте пароль для «Почта» (обычный пароль от аккаунта для SMTP не подойдёт).")
para("На сервере:")
code(["cd /var/www/era-coffee", "nano .env.local"])
para("Впишите (значения — пример для Яндекс 360):")
code([
    "SMTP_HOST=smtp.yandex.ru",
    "SMTP_PORT=465",
    "SMTP_USER=info@eracoffee.ru",
    "SMTP_PASS=пароль_приложения",
    "SMTP_FROM=ERA Coffee <info@eracoffee.ru>",
    "LEAD_EMAIL_TO=info@eracoffee.ru",
])
note("Для почты Reg.ru параметры аналогичны — host/порт берутся из панели почты Reg.ru.")
para("Пересоберите и перезапустите:")
code(["npm run build && pm2 restart era-coffee"])
para("Отправьте тестовую заявку с сайта — письмо должно прийти на info@eracoffee.ru. Если нет — смотрите pm2 logs era-coffee (строка [lead] ... и текст ошибки).")

h1("Шаг 12. Аналитика (опционально)")
para("Чтобы подключить Яндекс.Метрику и/или Google Analytics, впишите ID в .env.local:")
code([
    "NEXT_PUBLIC_YANDEX_METRIKA_ID=ваш_номер_счётчика",
    "NEXT_PUBLIC_GA_ID=G-XXXXXXXXXX",
])
para("Затем npm run build && pm2 restart era-coffee. Если переменные не заданы — счётчики просто не подключаются.")

h1("Шаг 13. Корпоративная почта info@eracoffee.ru")
para("Вариант А — почта от Reg.ru (платно, но просто): Reg.ru → Почта → Подключить для домена eracoffee.ru, далее по мастеру.")
para("Вариант Б — Яндекс 360 для бизнеса (есть бесплатный тариф):")
numbered("360.yandex.ru/business → добавить домен eracoffee.ru.")
numbered("Яндекс выдаст MX, TXT (SPF), DKIM записи — добавьте их в DNS-зону Reg.ru (тот же раздел, что в Шаге 8). Типичный MX:")
table(["Тип", "Имя", "Значение", "Приоритет"], [
    ["MX", "@", "mx.yandex.net.", "10"],
    ["TXT", "@", "v=spf1 redirect=_spf.yandex.net", "—"],
])
numbered("Создайте ящик info@eracoffee.ru в панели Яндекс 360.")
note("MX-записи на работу самого сайта не влияют — это отдельный канал для почты.")

h1("Обновление сайта в будущем")
para("Когда в GitHub появляются изменения, на сервере достаточно:")
code([
    "cd /var/www/era-coffee",
    "git pull",
    "npm ci",
    "npm run build",
    "pm2 restart era-coffee",
])

h1("Финальный чеклист")
for item in [
    "Домен eracoffee.ru оплачен на Reg.ru",
    "VPS (Ubuntu 22.04, 2+ ГБ RAM) создан, есть IP",
    "Node.js 20, nginx, PM2 установлены",
    "Репозиторий склонирован, npm run build прошёл",
    "PM2 держит era-coffee (порт 3000), pm2 save сделан",
    "DNS A-записи @ и www → IP VPS",
    "nginx-конфиг активен, nginx -t без ошибок",
    "SSL выпущен (https://eracoffee.ru открывается с замком)",
    "(для лидов) SMTP настроен в .env.local, формы шлют на info@eracoffee.ru",
    "(опц.) Метрика/GA ID прописаны",
    "(опц.) Почта info@eracoffee.ru настроена",
]:
    bullet(item, checklist=True)

h1("Если что-то не работает")
table(["Симптом", "Что проверить"], [
    ["502 Bad Gateway", "Запущен ли PM2: pm2 status. Жив ли порт: curl -I localhost:3000"],
    ["Домен не открывается", "DNS ещё не распространился / правильный ли IP в A-записи"],
    ["SSL не выпускается", "Домен уже должен указывать на сервер (Шаг 8 до Шага 10)"],
    ["Сборка падает (Killed)", "Мало RAM — добавьте swap (Шаг 6)"],
    ["Формы не приходят", "Проверьте SMTP_* в .env.local и pm2 logs era-coffee (Шаг 11)"],
])

doc.save(OUT)
print("✓ Wrote", OUT)
